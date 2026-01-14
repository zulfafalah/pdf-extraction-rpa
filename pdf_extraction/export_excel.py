# extract_excel.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from io import BytesIO


def collect_headers_from_json(extracted_results):
    header_fields = []
    item_fields = set()

    for _, result in extracted_results.items():
        if not result.get('success'):
            continue

        data = result.get('extracted_data', {})

        for k in data:
            if k != 'items' and k not in header_fields:
                header_fields.append(k)

        for item in data.get('items', []):
            item_fields.update(item.keys())

    return header_fields + sorted(item_fields)


def create_simple_export(extracted_results):
    headers = collect_headers_from_json(extracted_results)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extraction Results"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="366092")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # Header
    for col, field in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=field.upper().replace('_', ' '))
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        ws.column_dimensions[cell.column_letter].width = 25

    row = 2
    for _, result in extracted_results.items():
        if not result.get('success'):
            continue

        data = result.get('extracted_data', {})
        header_values = {k: v for k, v in data.items() if k != 'items'}
        items = data.get('items', [])

        if items:
            for item in items:
                for col, field in enumerate(headers, 1):
                    ws.cell(
                        row=row,
                        column=col,
                        value=item.get(field, header_values.get(field, ''))
                    ).border = border
                row += 1
        else:
            for col, field in enumerate(headers, 1):
                ws.cell(row=row, column=col, value=header_values.get(field, '')).border = border
            row += 1

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
